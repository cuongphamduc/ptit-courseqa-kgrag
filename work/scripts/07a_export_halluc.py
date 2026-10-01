"""Stage 7a: manual hallucination annotation files (primary SLM, PTIT test).

- 150 PTIT test questions, stratified by type (largest remainder over the test type distribution), seed 42;
  the same 150 questions for S0, S2, S4, S5 -> 600 answers.
- system names hidden (random codes), rows shuffled (seed 42). The key is written to work/halluc/key.json
  (not in ./annotation, annotators must not open it).
- halluc_A.xlsx: all 600 rows; halluc_B.xlsx: 20% of the rows (120), stratified by system x type.
Columns shown: item, question, gold answer + evidence (unanswerable: none), the context the system saw
(S0: none), the system answer. Label columns: label (faithful / partially_hallucinated / hallucinated /
correct_refusal / wrong_refusal) and note.
"""
import collections
import json
import os
import random
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from rag import data, llm  # noqa: E402

SEED, N_Q, B_FRAC = 42, 150, 0.20
SYSTEMS = ["S0", "S2", "S4", "S5"]
RES = os.path.join(data.ROOT, "results")
OUT = os.path.join(data.ROOT, "annotation", "halluc")
KEY = os.path.join(data.WORK, "halluc", "key.json")
LABELS = ["faithful", "partially_hallucinated", "hallucinated", "correct_refusal", "wrong_refusal"]


def allocate(total, weights):
    s = sum(weights.values())
    exact = {k: total * w / s for k, w in weights.items()}
    alloc = {k: int(x) for k, x in exact.items()}
    for k in sorted(exact, key=lambda k: (-(exact[k] - int(exact[k])), k))[:total - sum(alloc.values())]:
        alloc[k] += 1
    return alloc


def sample_questions(rng, n=N_Q, exclude=()):
    """stratified by type (largest remainder), deterministic given rng state"""
    by_type = collections.defaultdict(list)
    for q in data.questions("ptit", "test"):
        if q["id"] not in exclude:
            by_type[q["type"]].append(q)
    alloc = allocate(n, {t: len(v) for t, v in by_type.items()})
    sample = []
    for t in sorted(by_type):
        sample += rng.sample(sorted(by_type[t], key=lambda q: q["id"]), alloc[t])
    return sample


def main():
    rng = random.Random(SEED)
    raw = {json.loads(l)["id"]: json.loads(l) for l in open(os.path.join(data.BENCH, "ptit_test.jsonl"), encoding="utf-8")}
    chunks = {c["chunk_id"]: c for c in data.corpus("ptit")}
    sample = sample_questions(rng)
    preds, rets = {}, {}
    for s in SYSTEMS:
        preds[s] = {json.loads(l)["id"]: json.loads(l) for l in open(os.path.join(RES, "predictions", "ptit_test", f"{s}__{llm.PRIMARY}.jsonl"), encoding="utf-8")}
        if s != "S0":
            rets[s] = {json.loads(l)["id"]: json.loads(l) for l in open(os.path.join(RES, "retrieval", "ptit_test", f"{s}.jsonl"), encoding="utf-8")}
    codes = [f"H{c}" for c in rng.sample(range(10, 100), len(SYSTEMS))]
    sys_code = dict(zip(SYSTEMS, codes))
    rows = []
    for q in sample:
        r = raw[q["id"]]
        gold = "(câu hỏi không có đáp án trong tài liệu)" if q["unanswerable"] else r["answer"]
        ev = "\n\n".join(f"[{e['chunk_id']} | trang PDF {e['page']}] {e['span']}" for e in r["evidence"])
        for s in SYSTEMS:
            if s == "S0":
                ctx = "(hệ thống này không dùng tài liệu)"
            else:
                rr = rets[s][q["id"]]
                texts = rr.get("texts") or [chunks[c]["text"] for c in rr["chunk_ids"]]
                ctx = "\n\n".join(f"[{i + 1}] {t}" for i, t in enumerate(texts))
                if rr.get("facts"):
                    ctx += "\n\nQUAN HỆ:\n" + "\n".join(f"- {f}" for f in rr["facts"])
            rows.append({"qid": q["id"], "system": s, "type": q["type"], "question": q["question"], "gold_answer": gold,
                         "gold_evidence": ev, "context": ctx, "answer": preds[s][q["id"]]["answer"]})
    rng.shuffle(rows)
    for i, r in enumerate(rows, 1):
        r["item"] = f"HAL-{i:03d}"
    cells = collections.defaultdict(list)
    for r in rows:
        cells[(r["system"], r["type"])].append(r)
    b_alloc = allocate(round(len(rows) * B_FRAC), {k: len(v) for k, v in cells.items()})
    b_items = set()
    for k in sorted(cells):
        b_items |= {r["item"] for r in rng.sample(cells[k], b_alloc[k])}

    os.makedirs(OUT, exist_ok=True)
    os.makedirs(os.path.dirname(KEY), exist_ok=True)
    json.dump({"system_codes": sys_code, "items": {r["item"]: {"qid": r["qid"], "system": r["system"], "type": r["type"]} for r in rows},
               "b_items": sorted(b_items), "seed": SEED}, open(KEY, "w"), indent=1)
    cols = [("item", 10), ("question", 50), ("gold_answer", 40), ("gold_evidence", 60), ("context", 90), ("answer", 50),
            ("label", 22), ("note", 40)]
    for name, subset in (("halluc_A.xlsx", rows), ("halluc_B.xlsx", [r for r in rows if r["item"] in b_items])):
        wb = Workbook()
        ws = wb.active
        ws.title = "answers"
        ws.append([c for c, _ in cols])
        for j, (c, w) in enumerate(cols, 1):
            cell = ws.cell(row=1, column=j)
            cell.font = Font(bold=True)
            cell.fill = PatternFill("solid", fgColor="FFF2CC" if c in ("label", "note") else "DDDDDD")
            ws.column_dimensions[cell.column_letter].width = w
        for r in subset:
            ws.append([r["item"], r["question"], r["gold_answer"], r["gold_evidence"], r["context"], r["answer"], "", ""])
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        dv = DataValidation(type="list", formula1='"' + ",".join(LABELS) + '"', allow_blank=True, showErrorMessage=True)
        ws.add_data_validation(dv)
        dv.add(f"G2:G{len(subset) + 1}")
        ws.freeze_panes = "B2"
        rd = wb.create_sheet("README", 0)
        for line in ["Gán nhãn ảo giác – câu trả lời của 4 hệ thống ẩn danh (cùng một SLM), thứ tự đã xáo trộn.",
                     "Chỉ điền cột label và note. Định nghĩa nhãn: annotation/GUIDELINES.md, mục 'Gán nhãn ảo giác'.",
                     "label: " + " | ".join(LABELS),
                     "Không suy đoán hệ thống nào tạo câu trả lời. Không sửa cột nền xám."]:
            rd.append([line])
        rd.column_dimensions["A"].width = 140
        wb.active = 1
        wb.save(os.path.join(OUT, name))
    print({"A": len(rows), "B": len(b_items), "types": dict(collections.Counter(q["type"] for q in sample))})


if __name__ == "__main__":
    main()
