"""Stage 3e: export annotation workbooks (round 1) and draft statistics.

Output: annotation/round1/annotation_A.xlsx   all draft questions (test + dev)
        annotation/round1/annotation_B.xlsx   stratified 20% of TEST (type x book), seed 42, for agreement
        annotation/round1/overlap_ids.txt     IDs present in both A and B
        work/03_draft_stats.md
"""
import collections
import json
import os
import random

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DRAFT = os.path.join(ROOT, "work", "qgen", "merged", "ptit_draft.jsonl")
OUT = os.path.join(ROOT, "annotation", "round1")
SEED, B_FRACTION = 42, 0.20
BOOKS = json.load(open(os.path.join(ROOT, "work", "corpus", "books.json"), encoding="utf-8"))
CHUNKS = {json.loads(l)["chunk_id"]: json.loads(l) for l in open(os.path.join(ROOT, "work", "corpus", "all_chunks.jsonl"), encoding="utf-8")}

INFO_COLS = [("id", 16), ("split", 6), ("book_id", 11), ("type", 13), ("cross_type", 8), ("question", 60),
             ("proposed_answer", 40), ("evidence", 70), ("pdf_files", 30), ("reasoning_chain", 50),
             ("unanswerable_reason", 40), ("auto_check", 30), ("similar_questions", 18)]
ANN_COLS = [("label", 10), ("edited_question", 40), ("edited_answer", 30), ("type_correct", 10),
            ("evidence_correct", 10), ("note", 40)]


def fmt_evidence(q):
    lines = []
    for e in q["evidence"]:
        lines.append(f"[{e['chunk_id']} | {e['book_id']} | trang PDF {e['page']} | {e.get('section') or e.get('chapter')}]\n{e['span']}")
    return "\n\n".join(lines)


def fmt_auto(q):
    a = q.get("auto_check")
    if not a:
        return ""
    return f"{a['verdict']}" + (f" ({a['chunk_id']}): {a['note']}" if a.get("chunk_id") else (f": {a['note']}" if a.get("note") else ""))


def write_book(path, qs, title):
    wb = Workbook()
    ws = wb.active
    ws.title = "questions"
    head = [c for c, _ in INFO_COLS] + [c for c, _ in ANN_COLS]
    ws.append(head)
    fill_info = PatternFill("solid", fgColor="DDDDDD")
    fill_ann = PatternFill("solid", fgColor="FFF2CC")
    for j, (c, w) in enumerate(INFO_COLS + ANN_COLS, 1):
        cell = ws.cell(row=1, column=j)
        cell.font = Font(bold=True)
        cell.fill = fill_ann if j > len(INFO_COLS) else fill_info
        ws.column_dimensions[cell.column_letter].width = w
    for q in qs:
        files = sorted({f"dataset/PTIT/{BOOKS[e['book_id']]}" for e in q["evidence"]})
        ws.append([q["id"], q["split"], q["book_id"], q["type"], q.get("cross_type") or "", q["question"],
                   q["answer"] if q["answer"] is not None else "(không có đáp án)", fmt_evidence(q),
                   "\n".join(files), q.get("reasoning_chain") or "", q.get("unanswerable_reason") or "",
                   fmt_auto(q), ", ".join(q.get("similar_to") or [])] + [""] * len(ANN_COLS))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    ws.freeze_panes = "G2"
    n = len(qs) + 1
    col = {c: i + 1 for i, c in enumerate(head)}
    from openpyxl.utils import get_column_letter as L
    for name, opts in [("label", '"accept,edit,reject"'), ("type_correct", '"yes,no"'), ("evidence_correct", '"yes,no"')]:
        dv = DataValidation(type="list", formula1=opts, allow_blank=True, showErrorMessage=True)
        ws.add_data_validation(dv)
        dv.add(f"{L(col[name])}2:{L(col[name])}{n}")
    ws.auto_filter.ref = f"A1:{L(len(head))}{n}"

    # lookup sheet with full chunk text of every cited chunk
    cs = wb.create_sheet("chunks")
    cs.append(["chunk_id", "book_id", "pdf_file", "pages", "chapter", "section", "text"])
    for c in ("A", "B", "C", "D", "E", "F", "G"):
        cs.column_dimensions[c].width = {"G": 120, "C": 40}.get(c, 18)
    cited = sorted({e["chunk_id"] for q in qs for e in q["evidence"]})
    for cid in cited:
        ch = CHUNKS[cid]
        cs.append([cid, ch["book_id"], f"dataset/PTIT/{BOOKS[ch['book_id']]}", f"{ch['page_start']}-{ch['page_end']}",
                   ch["chapter"], ch["section"], ch["text"]])
    for row in cs.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    cs.freeze_panes = "B2"

    rd = wb.create_sheet("README")
    for line in [title, "",
                 "Chỉ điền các cột nền vàng: label, edited_question, edited_answer, type_correct, evidence_correct, note.",
                 "label: accept = giữ nguyên | edit = giữ sau khi sửa (ghi edited_question và/hoặc edited_answer) | reject = loại.",
                 "type_correct: loại câu hỏi (cột type/cross_type) có đúng không (yes/no).",
                 "evidence_correct: bằng chứng (chunk + span + trang) có đủ và đúng để trả lời không (yes/no). Với unanswerable: chunk gây nhầm có hợp lý và tài liệu đúng là không có đáp án không.",
                 "similar_questions: các câu có cách hỏi giống câu này (TF-IDF ký tự, cos >= 0,6). Nếu trùng ý với câu đã accept, hãy reject câu sau và ghi ID vào note.",
                 "auto_check (chỉ với unanswerable): kết quả kiểm tra tự động trên 11 giáo trình; 'answerable' nghĩa là máy nghi có đáp án ở chunk được nêu. Hãy kiểm tra kỹ các dòng này.",
                 "Sheet 'chunks' chứa toàn văn các chunk được dẫn; mở PDF ở cột pdf_files, đúng số trang PDF (không phải số in ở chân trang).",
                 "Hướng dẫn đầy đủ: annotation/GUIDELINES.md. Không đổi thứ tự hay xóa dòng; không sửa các cột nền xám."]:
        rd.append([line])
    rd.column_dimensions["A"].width = 150
    wb.move_sheet("README", offset=-2)
    wb.active = 1
    wb.save(path)


def stratified_b(qs):
    rng = random.Random(SEED)
    test = [q for q in qs if q["split"] == "test"]
    cells = collections.defaultdict(list)
    for q in test:
        cells[(q["type"], q["book_id"])].append(q)
    target = round(len(test) * B_FRACTION)
    exact = {k: len(v) * B_FRACTION for k, v in cells.items()}
    alloc = {k: int(x) for k, x in exact.items()}
    rest = target - sum(alloc.values())
    for k in sorted(exact, key=lambda k: -(exact[k] - alloc[k]))[:rest]:
        alloc[k] += 1
    chosen = []
    for k in sorted(cells):
        chosen += rng.sample(cells[k], alloc[k])
    return sorted(chosen, key=lambda q: q["id"])


SIM_THRESHOLD = 0.60  # char 3-5 gram TF-IDF cosine; shown to annotators as a duplicate hint only


def add_similar(qs):
    from sklearn.feature_extraction.text import TfidfVectorizer
    from sklearn.metrics.pairwise import cosine_similarity
    for split in ("test", "dev"):
        s = [q for q in qs if q["split"] == split]
        sim = cosine_similarity(TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5)).fit_transform(
            [q["question"].lower() for q in s]))
        for i, q in enumerate(s):
            q["similar_to"] = [s[j]["id"] for j in range(len(s)) if j != i and sim[i, j] >= SIM_THRESHOLD]


def main():
    os.makedirs(OUT, exist_ok=True)
    qs = [json.loads(l) for l in open(DRAFT, encoding="utf-8")]
    add_similar(qs)
    write_book(os.path.join(OUT, "annotation_A.xlsx"), qs, "PTIT benchmark – gán nhãn vòng 1 – Người gán nhãn A (toàn bộ câu hỏi)")
    b = stratified_b(qs)
    write_book(os.path.join(OUT, "annotation_B.xlsx"), b, "PTIT benchmark – gán nhãn vòng 1 – Người gán nhãn B (20% câu test, phân tầng theo loại × giáo trình, seed 42)")
    open(os.path.join(OUT, "overlap_ids.txt"), "w").write("\n".join(q["id"] for q in b) + "\n")

    # stats report
    types = ["single", "multi_intra", "multi_cross", "unanswerable"]
    def table(split):
        sub = [q for q in qs if q["split"] == split]
        rows = ["| Giáo trình | " + " | ".join(types) + " | Tổng |", "|---|" + "---:|" * (len(types) + 1)]
        tot = collections.Counter()
        for bk in BOOKS:
            c = collections.Counter(q["type"] for q in sub if q["book_id"] == bk)
            tot.update(c)
            rows.append(f"| {bk} | " + " | ".join(str(c[t]) for t in types) + f" | {sum(c.values())} |")
        n = sum(tot.values())
        rows.append("| **Tổng** | " + " | ".join(f"**{tot[t]}** ({tot[t] / max(1, n):.1%})" for t in types) + f" | **{n}** |")
        return "\n".join(rows), n
    tt, nt = table("test")
    td, nd = table("dev")
    cross = collections.Counter(f"{q['split']}/{q.get('cross_type')}" for q in qs if q["type"] == "multi_cross")
    auto = collections.Counter(q["auto_check"]["verdict"] for q in qs if q.get("auto_check"))
    rep = json.load(open(os.path.join(ROOT, "work", "qgen", "merged", "merge_report.json"), encoding="utf-8"))
    bcnt = collections.Counter(q["type"] for q in b)
    md = f"""# Giai đoạn 3: Thống kê bộ câu hỏi nháp (vòng 1)

Mô hình sinh: `claude-opus-5-5` (subagent của Claude Code), khác với các SLM được đánh giá.
Script: `work/scripts/03a_plan_qgen.py` → sinh bằng agent (`work/qgen/INSTRUCTIONS.md`) → `03b_validate_questions.py` → `03c_merge_questions.py` → `03d_unanswerable_check.py` → `03e_export_annotation.py`.
Cấu hình: `results/configs/qgen_config.json`.

- Câu nháp giữ lại: **{rep['kept']}** (test {nt}, dev {nd}).
- Số câu bị loại khi gộp: {rep['dropped']}. Lý do: {dict(rep['dropped_reasons'])}.
- Câu `multi_cross` chia theo `cross_type`: {dict(cross)}.
- Kiểm tra tự động câu `unanswerable` trên toàn kho (BM25 top-8, agent kiểm định): {dict(auto)}.

## Test (mục tiêu sau duyệt: 1.000 câu; 35/30/20/15%)

{tt}

## Dev (mục tiêu sau duyệt: 200 câu)

{td}

## File gán nhãn

- `annotation/round1/annotation_A.xlsx`: {len(qs)} câu (toàn bộ test + dev).
- `annotation/round1/annotation_B.xlsx`: {len(b)} câu = 20% câu test, lấy mẫu phân tầng theo loại × giáo trình, seed {SEED}. Theo loại: {dict(bcnt)}. Danh sách ID ở `annotation/round1/overlap_ids.txt`.
- Dev và test không dùng chung chunk bằng chứng (vùng chunk chia trước, seed 42, `work/qgen/plan.json`).
"""
    open(os.path.join(ROOT, "work", "03_draft_stats.md"), "w", encoding="utf-8").write(md)
    print(md)


if __name__ == "__main__":
    main()
