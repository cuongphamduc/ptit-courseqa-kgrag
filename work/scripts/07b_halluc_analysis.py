"""Stage 7b (after GATE 3): manual hallucination labels -> agreement, decoded per-system rates, validation of the
automatic metrics.

  agreement   Cohen's kappa A-B (5 labels; binary hallucination = hallucinated|partially_hallucinated)
  rates       per system (S0, S2, S4, S5; qwen3.5-4b; 150 PTIT test questions): shares of each label,
              hallucination rate over all answers and over non-refusal answers, 95% bootstrap CI (10,000, seed 42)
  tests       exact McNemar S5 vs S0/S2/S4 on binary hallucination (paired by question), Holm
  validation  Spearman between the manual score (faithful 1, partial 0.5, hallucinated 0) and the judge faithfulness
              score (RAG systems, non-refusal) and token F1 (answerable, non-refusal); agreement of the manual refusal
              labels with the automatic refusal detector; label consistency with the gold answerability
Output: results/halluc/manual_halluc.json
"""
import collections
import json
import os
import sys

import numpy as np
from openpyxl import load_workbook
from scipy.stats import binomtest, spearmanr
from sklearn.metrics import cohen_kappa_score

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from rag import data  # noqa: E402
from rag.text import f1_score, is_refusal  # noqa: E402

ROOT, RES = data.ROOT, os.path.join(data.ROOT, "results")
KEY = json.load(open(os.path.join(data.WORK, "halluc", "key.json")))
LABELS = ["faithful", "partially_hallucinated", "hallucinated", "correct_refusal", "wrong_refusal"]
HALLUC = {"partially_hallucinated", "hallucinated"}
SCORE = {"faithful": 1.0, "partially_hallucinated": 0.5, "hallucinated": 0.0}
SYSTEMS = ["S0", "S2", "S4", "S5"]


def read(name):
    ws = load_workbook(os.path.join(ROOT, "annotation", "halluc", name), read_only=True)["answers"]
    it = ws.iter_rows(values_only=True)
    h = list(next(it))
    return {r[0]: dict(zip(h, r)) for r in it if r[0]}


def boot_ci(x, rng, b=10000):
    x = np.asarray(x, dtype=float)
    m = x[rng.integers(0, len(x), size=(b, len(x)))].mean(axis=1)
    return [round(float(np.percentile(m, 2.5)), 4), round(float(np.percentile(m, 97.5)), 4)]


def holm(ps):
    order = sorted(range(len(ps)), key=lambda i: ps[i])
    adj, run = [0.0] * len(ps), 0.0
    for r, i in enumerate(order):
        run = max(run, min(1.0, (len(ps) - r) * ps[i]))
        adj[i] = run
    return adj


def main():
    A, B = read("halluc_A.xlsx"), read("halluc_B.xlsx")
    lab = lambda r: (r["label"] or "").strip()
    assert all(lab(r) in LABELS for r in A.values()) and all(lab(r) in LABELS for r in B.values())
    ids = sorted(B)
    out = {"n_A": len(A), "n_B": len(B),
           "provenance": {"notes_filled_A": sum(1 for r in A.values() if (r["note"] or "").strip()),
                          "notes_filled_B": sum(1 for r in B.values() if (r["note"] or "").strip()),
                          "B_notes_identical_to_A": sum(1 for i in B if (B[i]["note"] or "").strip() and " ".join(str(B[i]["note"]).split()) == " ".join(str(A[i]["note"] or "").split())),
                          "rows_both_flagged_with_notes": sum(1 for i in B if (B[i]["note"] or "").strip() and (A[i]["note"] or "").strip())}}
    a, b = [lab(A[i]) for i in ids], [lab(B[i]) for i in ids]
    ab, bb = [x in HALLUC for x in a], [x in HALLUC for x in b]
    out["agreement"] = {"kappa_5class": round(float(cohen_kappa_score(a, b)), 4), "raw_5class": round(float(np.mean([x == y for x, y in zip(a, b)])), 4),
                        "kappa_binary_halluc": round(float(cohen_kappa_score(ab, bb)), 4), "raw_binary": round(float(np.mean([x == y for x, y in zip(ab, bb)])), 4),
                        "confusion": {f"{x}->{y}": c for (x, y), c in collections.Counter(zip(a, b)).items() if x != y}}

    qs = {q["id"]: q for q in data.questions("ptit", "test")}
    preds = {s: {json.loads(l)["id"]: json.loads(l) for l in open(os.path.join(RES, "predictions", "ptit_test", f"{s}__qwen3.5-4b.jsonl"), encoding="utf-8")} for s in SYSTEMS}
    judge = {}
    for s in SYSTEMS[1:]:
        p = os.path.join(RES, "judge", "ptit_test", f"{s}__qwen3.5-4b.jsonl")
        judge[s] = {json.loads(l)["id"]: json.loads(l) for l in open(p, encoding="utf-8")}
    rows = []
    for item, r in A.items():
        k = KEY["items"][item]
        q = qs[k["qid"]]
        ans = preds[k["system"]][k["qid"]]["answer"]
        rows.append({"item": item, "qid": k["qid"], "system": k["system"], "type": k["type"], "label": lab(r),
                     "unanswerable": q["unanswerable"], "auto_refusal": is_refusal(ans),
                     "f1": None if q["unanswerable"] else f1_score(ans, q["golds"]),
                     "judge": (judge.get(k["system"], {}).get(k["qid"]) or {}).get("score")})

    rng = np.random.default_rng(42)
    rates = {}
    for s in SYSTEMS:
        rs = [r for r in rows if r["system"] == s]
        c = collections.Counter(r["label"] for r in rs)
        h_all = [r["label"] in HALLUC for r in rs]
        answered = [r for r in rs if r["label"] not in ("correct_refusal", "wrong_refusal")]
        h_ans = [r["label"] in HALLUC for r in answered]
        una = [r for r in rs if r["unanswerable"]]
        rates[s] = {"n": len(rs), "labels": {l: c[l] for l in LABELS},
                    "halluc_rate_all": round(float(np.mean(h_all)), 4), "halluc_rate_all_ci95": boot_ci(h_all, rng),
                    "strict_halluc_rate_all": round(float(np.mean([r["label"] == "hallucinated" for r in rs])), 4),
                    "halluc_rate_answered": round(float(np.mean(h_ans)), 4) if h_ans else None,
                    "halluc_rate_answered_ci95": boot_ci(h_ans, rng) if h_ans else None, "n_answered": len(answered),
                    "faithful_rate_all": round(c["faithful"] / len(rs), 4),
                    "unanswerable_halluc_rate": round(float(np.mean([r["label"] in HALLUC for r in una])), 4) if una else None,
                    "by_type_halluc_rate": {t: round(float(np.mean([r["label"] in HALLUC for r in rs if r["type"] == t])), 4)
                                            for t in sorted({r["type"] for r in rs})}}
    out["rates"] = rates

    by = {(r["system"], r["qid"]): r["label"] in HALLUC for r in rows}
    qids = sorted({r["qid"] for r in rows})
    tests = {}
    for base in ("S0", "S2", "S4"):
        b_ = sum(1 for q in qids if not by[("S5", q)] and by[(base, q)])   # S5 better (no hallucination) only
        c_ = sum(1 for q in qids if by[("S5", q)] and not by[(base, q)])
        tests[base] = {"s5_better_only": b_, "baseline_better_only": c_, "p": float(binomtest(b_, b_ + c_, 0.5).pvalue) if b_ + c_ else 1.0}
    for n, p in zip(list(tests), holm([t["p"] for t in tests.values()])):
        tests[n]["p_holm"] = round(p, 6)
    out["mcnemar_s5_vs"] = tests

    ans_rows = [r for r in rows if r["label"] in SCORE]
    jr = [r for r in ans_rows if r["judge"] is not None]
    fr = [r for r in ans_rows if r["f1"] is not None]
    rho_j = spearmanr([SCORE[r["label"]] for r in jr], [r["judge"] for r in jr])
    rho_f = spearmanr([SCORE[r["label"]] for r in fr], [r["f1"] for r in fr])
    man_ref = [r["label"] in ("correct_refusal", "wrong_refusal") for r in rows]
    auto_ref = [r["auto_refusal"] for r in rows]
    out["validation"] = {
        "spearman_manual_vs_judge": {"rho": round(float(rho_j.statistic), 4), "p": float(rho_j.pvalue), "n": len(jr)},
        "spearman_manual_vs_f1": {"rho": round(float(rho_f.statistic), 4), "p": float(rho_f.pvalue), "n": len(fr)},
        "refusal_detector_vs_manual": {"agreement": round(float(np.mean([x == y for x, y in zip(man_ref, auto_ref)])), 4),
                                       "kappa": round(float(cohen_kappa_score(man_ref, auto_ref)), 4),
                                       "manual_refusal_auto_answer": sum(1 for x, y in zip(man_ref, auto_ref) if x and not y),
                                       "manual_answer_auto_refusal": sum(1 for x, y in zip(man_ref, auto_ref) if y and not x)},
        "label_consistency": {"correct_refusal_on_answerable": sum(1 for r in rows if r["label"] == "correct_refusal" and not r["unanswerable"]),
                              "wrong_refusal_on_unanswerable": sum(1 for r in rows if r["label"] == "wrong_refusal" and r["unanswerable"]),
                              "faithful_on_unanswerable": sum(1 for r in rows if r["label"] == "faithful" and r["unanswerable"])}}
    os.makedirs(os.path.join(RES, "halluc"), exist_ok=True)
    json.dump(out, open(os.path.join(RES, "halluc", "manual_halluc.json"), "w"), indent=1, ensure_ascii=False)
    with open(os.path.join(RES, "halluc", "manual_labels_decoded.jsonl"), "w", encoding="utf-8") as f:
        for r in rows:
            f.write(json.dumps(r, ensure_ascii=False) + "\n")
    print(json.dumps(out, indent=1, ensure_ascii=False))


if __name__ == "__main__":
    main()
