"""Stage 3g (GATE 2 processing): read round-1 annotation, compute agreement, apply edits, finalise the benchmark.

Steps (AGENTS.MD, GATE 2):
  1. every row must have `label`;
  2. Cohen's kappa on the A-B overlap for label / type_correct / evidence_correct (+ raw agreement);
  3. apply edits, drop `reject`; when type_correct=no and the note names the single sufficient chunk
     ("Loại đúng: single ... <chunk_id>"), the question becomes `single` with that chunk as its only evidence;
     three `evidence_correct=no` edits get the missing sentences added as extra spans (EVIDENCE_ADD);
  4. stratified down-sampling of the accepted pool (seed 42) to the planned sizes: test 1,000 and dev 200,
     type ratio 35/30/20/15 %, allocation across books proportional to availability.

Output: work/benchmark/{ptit_test.jsonl, ptit_dev.jsonl, ptit_accepted_pool.jsonl}
        results/annotation_round1_agreement.json, results/benchmark_stats.json
"""
import collections
import importlib.util
import json
import os
import random
import re

from openpyxl import load_workbook
from sklearn.metrics import cohen_kappa_score

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DRAFT = os.path.join(ROOT, "work", "qgen", "merged", "ptit_draft.jsonl")
ANN = os.path.join(ROOT, "annotation", "round1")
OUT = os.path.join(ROOT, "work", "benchmark")
RES = os.path.join(ROOT, "results")
SEED = 42
TARGET = {"test": 1000, "dev": 200}
RATIO = {"single": 0.35, "multi_intra": 0.30, "multi_cross": 0.20, "unanswerable": 0.15}
CHUNKS = {json.loads(l)["chunk_id"]: json.loads(l) for l in open(os.path.join(ROOT, "work", "corpus", "all_chunks.jsonl"), encoding="utf-8")}
BOOKS = json.load(open(os.path.join(ROOT, "work", "corpus", "books.json"), encoding="utf-8"))

spec = importlib.util.spec_from_file_location("fixpages", os.path.join(ROOT, "work", "scripts", "03f_fix_pages.py"))
fixpages = importlib.util.module_from_spec(spec)
spec.loader.exec_module(fixpages)

# sentences the annotators said were missing from the spans (notes of these rows); copied verbatim from the chunk text
EVIDENCE_ADD = {
    "PTIT-DEV-0254": [("MKT-QTBH-0122", "- Nghiên cứu thị trường để hiểu rõ nhu cầu của khách hàng lớn."),
                      ("MKT-QTBH-0123", "- Thiết kế phối thức marketing hỗn hợp riêng biệt nhằm ưu đãi khách hàng lớn.")],
    "PTIT-TEST-0984": [("MKT-CRM-0060", "Cấp độ 3: bán hàng trực tuyến thông qua hệ thống thương mại điện tử.")],
    "PTIT-TEST-1264": [("MKT-QTBH-0091", "- Hoạt động kiểm tra, đánh giá phải phù hợp với công tác tổ chức và nhân sự trong lực lượng bán hàng "
                                         "(bao gồm cả quản trị bán hàng viên và nhân viên bán hàng)"),
                       ("MKT-QTBH-0091", "- Hệ thống kiểm tra, đánh giá cần phải khách quan: việc đánh giá phải được thực hiện một cách chính xác, "
                                         "đúng nguyên tắc, phải dựa vào những tiêu chuẩn rõ ràng.")],
}
ANN_COLS = ["label", "edited_question", "edited_answer", "type_correct", "evidence_correct", "note"]


def read_sheet(path):
    ws = load_workbook(path, read_only=True)["questions"]
    it = ws.iter_rows(values_only=True)
    head = list(next(it))
    return {r[0]: dict(zip(head, r)) for r in it if r[0]}


def ws(s):
    return re.sub(r"\s+", " ", s or "").strip()


def page_of(book, cid, span):
    c, s = CHUNKS[cid], fixpages.n(span)
    for key in (s[:30], s[len(s) // 2:len(s) // 2 + 30], s[-30:]):
        for p in range(max(1, c["page_start"] - 1), c["page_end"] + 2):
            if key and key in fixpages.page_text(book, p):
                return p
    return c["page_start"]


def agreement(A, B):
    ids = sorted(B)
    out = {"n_overlap": len(ids)}
    for col in ("label", "type_correct", "evidence_correct"):
        a, b = [A[i][col] for i in ids], [B[i][col] for i in ids]
        k = cohen_kappa_score(a, b) if len(set(a) | set(b)) > 1 else float("nan")
        out[col] = {"kappa": None if k != k else round(float(k), 4), "raw_agreement": round(sum(x == y for x, y in zip(a, b)) / len(ids), 4),
                    "dist_A": dict(collections.Counter(a)), "dist_B": dict(collections.Counter(b))}
    same = {c: sum(ws(str(A[i][c] or "")) == ws(str(B[i][c] or "")) for i in ids) for c in ANN_COLS}
    out["identical_cells_A_vs_B"] = same
    out["rows_identical_on_all_label_columns"] = sum(all(ws(str(A[i][c] or "")) == ws(str(B[i][c] or "")) for c in ANN_COLS[:5]) for i in ids)
    return out


def apply(q, r):
    q = json.loads(json.dumps(q))
    for k in ("tmp_id", "_source", "auto_warnings", "status", "similar_to"):
        q.pop(k, None)
    ann = {"label": r["label"], "type_correct": r["type_correct"], "evidence_correct": r["evidence_correct"], "note": ws(r["note"])}
    if r["label"] == "edit":
        if ws(r["edited_question"]):
            ann["original_question"] = q["question"]
            q["question"] = ws(r["edited_question"])
        if ws(r["edited_answer"]):
            ann["original_answer"] = q["answer"]
            q["answer"] = ws(r["edited_answer"])
        if r["type_correct"] == "no":
            m = re.search(r"Loại đúng:\s*single\b.*?([A-Z]+-[A-Z]+-\d{4})", ann["note"])
            assert m, (q["id"], ann["note"])
            cid = m.group(1)
            keep = [e for e in q["evidence"] if e["chunk_id"] == cid]
            assert keep, (q["id"], cid)
            ann.update(original_type=q["type"], original_cross_type=q.get("cross_type"), original_evidence=q["evidence"],
                       original_reasoning_chain=q.get("reasoning_chain"))
            q.update(type="single", cross_type=None, evidence=keep, book_id=CHUNKS[cid]["book_id"],
                     reasoning_chain=f"Đáp án nằm ở {cid} (loại câu sửa từ {ann['original_type']} khi gán nhãn).")
        for cid, span in EVIDENCE_ADD.get(q["id"], []):
            assert ws(span) in ws(CHUNKS[cid]["text"]), (q["id"], cid)
            c = CHUNKS[cid]
            q["evidence"].append({"chunk_id": cid, "book_id": c["book_id"], "page": page_of(c["book_id"], cid, span), "span": span,
                                  "chapter": c["chapter"], "section": c["section"]})
            ann["evidence_extended"] = True
    q["status"] = "final"
    q["annotation"] = ann
    return q


def allocate(total, weights):
    """largest-remainder allocation of `total` proportional to `weights` (dict), capped by the weights themselves."""
    s = sum(weights.values())
    exact = {k: total * w / s for k, w in weights.items()}
    alloc = {k: min(int(x), weights[k]) for k, x in exact.items()}
    for k in sorted(exact, key=lambda k: -(exact[k] - int(exact[k]))):
        if sum(alloc.values()) >= total:
            break
        if alloc[k] < weights[k]:
            alloc[k] += 1
    return alloc


def sample(pool, split):
    rng = random.Random(SEED)
    items = [q for q in pool if q["split"] == split]
    n = TARGET[split]
    type_target = allocate(n, {t: RATIO[t] * 1000 for t in RATIO})
    chosen = []
    for t in RATIO:
        cells = collections.defaultdict(list)
        for q in items:
            if q["type"] == t:
                cells[q["book_id"]].append(q)
        assert sum(map(len, cells.values())) >= type_target[t], (split, t)
        alloc = allocate(type_target[t], {b: len(v) for b, v in cells.items()})
        for b in sorted(cells):
            chosen += rng.sample(sorted(cells[b], key=lambda q: q["id"]), alloc[b])
    return sorted(chosen, key=lambda q: q["id"])


def dist(qs):
    n = len(qs)
    t = collections.Counter(q["type"] for q in qs)
    return {"n": n, "type": {k: t[k] for k in RATIO}, "type_pct": {k: round(100 * t[k] / max(1, n), 2) for k in RATIO},
            "cross_type": dict(collections.Counter(q["cross_type"] for q in qs if q["type"] == "multi_cross")),
            "book": {b: sum(q["book_id"] == b for q in qs) for b in BOOKS},
            "book_x_type": {b: {k: sum(q["book_id"] == b and q["type"] == k for q in qs) for k in RATIO} for b in BOOKS}}


def describe(qs):
    """lengths in syllables (whitespace tokens) and evidence counts"""
    def mean(xs):
        return round(sum(xs) / max(1, len(xs)), 2)
    ans = [q for q in qs if q["answer"] is not None]
    return {"question_len_mean": mean([len(q["question"].split()) for q in qs]),
            "answer_len_mean": mean([len(q["answer"].split()) for q in ans]),
            "answer_len_median": sorted(len(q["answer"].split()) for q in ans)[len(ans) // 2],
            "evidence_items_mean": mean([len(q["evidence"]) for q in qs]),
            "distinct_evidence_chunks": len({e["chunk_id"] for q in qs for e in q["evidence"]}),
            "multi_hop_distinct_chunks_mean": mean([len({e["chunk_id"] for e in q["evidence"]}) for q in qs if q["type"].startswith("multi")]),
            "edited_by_annotators": sum(q["annotation"]["label"] == "edit" for q in qs),
            "type_changed_by_annotators": sum(bool(q["annotation"].get("original_type")) for q in qs)}


def main():
    os.makedirs(OUT, exist_ok=True)
    draft = {json.loads(l)["id"]: json.loads(l) for l in open(DRAFT, encoding="utf-8")}
    A = read_sheet(os.path.join(ANN, "annotation_A.xlsx"))
    B = read_sheet(os.path.join(ANN, "annotation_B.xlsx"))
    missing = [i for i in draft if i not in A or not A[i]["label"]] + [i for i in B if not B[i]["label"]]
    assert not missing, missing
    assert set(B) == set(open(os.path.join(ANN, "overlap_ids.txt")).read().split())

    agr = agreement(A, B)
    json.dump(agr, open(os.path.join(RES, "annotation_round1_agreement.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)

    pool = [apply(draft[i], A[i]) for i in sorted(draft) if A[i]["label"] != "reject"]
    rejected = {i: ws(A[i]["note"]) for i in sorted(draft) if A[i]["label"] == "reject"}
    with open(os.path.join(OUT, "ptit_accepted_pool.jsonl"), "w", encoding="utf-8") as f:
        for q in pool:
            f.write(json.dumps(q, ensure_ascii=False) + "\n")

    final = {s: sample(pool, s) for s in TARGET}
    for s, qs in final.items():
        with open(os.path.join(OUT, f"ptit_{s}.jsonl"), "w", encoding="utf-8") as f:
            for q in qs:
                f.write(json.dumps(q, ensure_ascii=False) + "\n")

    labels = collections.Counter(r["label"] for r in A.values())
    stats = {
        "seed": SEED,
        "draft": {s: dist([q for q in draft.values() if q["split"] == s]) for s in TARGET},
        "annotation_A_labels": dict(labels),
        "annotation_A_labels_by_split": {s: dict(collections.Counter(A[i]["label"] for i in draft if draft[i]["split"] == s)) for s in TARGET},
        "type_changed_to_single": sum(1 for q in pool if q["annotation"].get("original_type")),
        "type_changed_from": dict(collections.Counter(f"{q['annotation']['original_type']}/{q['annotation']['original_cross_type']}"
                                                      for q in pool if q["annotation"].get("original_type"))),
        "question_edited": sum(1 for q in pool if "original_question" in q["annotation"]),
        "answer_edited": sum(1 for q in pool if "original_answer" in q["annotation"]),
        "evidence_extended": sum(1 for q in pool if q["annotation"].get("evidence_extended")),
        "rejected": rejected,
        "accepted_pool": {s: dist([q for q in pool if q["split"] == s]) for s in TARGET},
        "final": {s: dist(qs) for s, qs in final.items()},
        "final_B_overlap_in_test": sum(1 for q in final["test"] if q["id"] in B),
        "edited_questions_naming_a_textbook": [q["id"] for q in pool if "original_question" in q["annotation"] and re.search(r"giáo trình|sách", q["question"], re.I)],
        "descriptive": {s: describe(qs) for s, qs in final.items()},
        "evidence_chunks_shared_test_dev": len({e["chunk_id"] for q in final["test"] for e in q["evidence"]}
                                               & {e["chunk_id"] for q in final["dev"] for e in q["evidence"]}),
    }
    json.dump(stats, open(os.path.join(RES, "benchmark_stats.json"), "w", encoding="utf-8"), ensure_ascii=False, indent=1)
    print(json.dumps({k: v for k, v in agr.items()}, ensure_ascii=False, indent=1))
    print(json.dumps({k: stats[k] for k in ("annotation_A_labels", "type_changed_to_single", "type_changed_from", "question_edited",
                                             "answer_edited", "evidence_extended", "final_B_overlap_in_test", "edited_questions_naming_a_textbook")},
                     ensure_ascii=False, indent=1))
    for s in TARGET:
        print(s, "pool", stats["accepted_pool"][s]["type"], stats["accepted_pool"][s]["type_pct"])
        print(s, "final", stats["final"][s]["type"], stats["final"][s]["type_pct"], stats["final"][s]["cross_type"])
        print(s, "books", stats["final"][s]["book"])


if __name__ == "__main__":
    main()
